"""
report.py
---------
Formats the ranked shortlist as a polished Excel workbook: bold header,
frozen header row, auto-sized columns, percentage/number formatting, and
a conditional colour scale on the final score so the best funds pop
visually. Falls back gracefully (prints a warning) if openpyxl isn't
installed rather than crashing the whole pipeline.
"""

import pandas as pd

PERCENT_COLS = {
    "cagr_1y", "cagr_3y", "cagr_5y", "volatility", "max_drawdown",
    "downside_deviation", "rolling_avg", "rolling_min", "rolling_std",
    "rolling_positive_pct", "sip_xirr",
    "benchmark_cagr", "benchmark_excess_return", "rolling_beat_benchmark_pct",
    "pq_coverage", "score_coverage",
    "downside_capture", "quartile_consistency", "q1_share", "max_drawdown_all", "jensen_alpha",
}
# portfolio_quality is on a 0-100 scale; the pq_* inputs are already in percent units.
RATIO_COLS = {"beta", "risk_return_ratio", "sortino_ratio", "composite_score", "portfolio_quality",
              "pq_top10_pct", "pq_top_sector_pct", "pq_largemid_pct", "pq_cash_pct",
              "final_score_0_100"}


def _is_points_col(name: str) -> bool:
    """score_<metric> (0-100 percentile score) and contrib_<metric> (points) columns."""
    return (name.startswith("contrib_") or name.startswith("score_")) and name not in (
        "score_coverage", "score_missing")


def export_to_excel(df: pd.DataFrame, path: str) -> bool:
    """Writes df to an .xlsx file with readable formatting. Returns True on
    success, False if openpyxl isn't available (caller should fall back to
    CSV-only in that case)."""
    try:
        import openpyxl  # noqa: F401
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  (openpyxl not installed -- skipping Excel export. `pip install openpyxl` to enable it.)")
        return False

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Shortlist")
        ws = writer.sheets["Shortlist"]

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A2"

        for col_idx, col_name in enumerate(df.columns, start=1):
            letter = get_column_letter(col_idx)
            max_len = max(
                [len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str).tolist()]
            )
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 32)

            if col_name in PERCENT_COLS:
                fmt = "0.00%"
            elif col_name in RATIO_COLS:
                fmt = "0.00"
            elif _is_points_col(col_name):
                fmt = "0.0"
            else:
                fmt = None
            if fmt:
                for row_idx in range(2, len(df) + 2):
                    ws.cell(row=row_idx, column=col_idx).number_format = fmt

        # colour the 0-100 final score when present (it is blank for INSUFFICIENT_DATA funds),
        # otherwise the composite score
        score_name = "final_score_0_100" if "final_score_0_100" in df.columns else (
            "composite_score" if "composite_score" in df.columns else None)
        if score_name:
            score_col_idx = list(df.columns).index(score_name) + 1
            letter = get_column_letter(score_col_idx)
            rng = f"{letter}2:{letter}{len(df) + 1}"
            rule = ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B",
            )
            ws.conditional_formatting.add(rng, rule)

    return True
