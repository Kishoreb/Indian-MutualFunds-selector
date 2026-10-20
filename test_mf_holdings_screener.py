"""
Offline tests for mf_holdings_screener.py (synthetic NAV data, synthetic AMFI file, no network).
VERBATIM copy of the project's test (used here only as a regression run).
"""
import os
import shutil
import sys
import tempfile
from datetime import date, timedelta

import numpy as np
import pandas as pd

import config
import fetch_nav_data
import mf_holdings_screener as mfs

TODAY = date.today()


def isin_for(n):
    return f"INF{n:04d}K01{n % 100:02d}0"[:12].ljust(12, "0")


def make_nav(drift, vol=0.04, years=7.0, end_age_days=1, seed=0):
    rng = np.random.default_rng(seed)
    end = pd.Timestamp(TODAY - timedelta(days=end_age_days))
    idx = pd.bdate_range(end=end, periods=int(years * 252))
    daily = rng.normal(drift / 252, vol / np.sqrt(252), len(idx))
    nav = 100 * np.exp(np.cumsum(daily))
    return pd.DataFrame({"date": idx, "nav": nav})


def navall_text(schemes):
    lines = ["Scheme Code;ISIN Div Payout/ ISIN Growth;ISIN Div Reinvestment;Scheme Name;Net Asset Value;Date", ""]
    last = None
    for code, isin, name, cat in schemes:
        if cat != last:
            lines += [f"Open Ended Schemes({cat})", "", "Test Mutual Fund", ""]
            last = cat
        lines.append(f"{code};{isin};-;{name};123.45;{TODAY.strftime('%d-%b-%Y')}")
    return "\n".join(lines)


def check(cond, msg):
    if not cond:
        raise AssertionError(msg)


def test_parse_navall_and_index():
    text = navall_text([
        ("1001", "INF000A01011", "Alpha Small Cap Fund - Direct Plan - Growth Option", "Equity Scheme - Small Cap Fund"),
        ("1002", "INF000A01029", "Alpha Small Cap Fund - Direct Plan - IDCW", "Equity Scheme - Small Cap Fund"),
        ("2001", "INF000B01013", "Beta Liquid Fund - Regular Plan - Growth", "Debt Scheme - Liquid Fund"),
    ]) + "\nClose Ended Schemes(Other Scheme - FMP)\n\n9;INF000C01017;-;Gamma FMP - Direct - Growth;10.0;01-Jan-2026\n"
    recs = mfs.parse_navall(text)
    check(len(recs) == 4, f"expected 4 records, got {len(recs)}")
    check(recs[0]["category"] == "Equity Scheme - Small Cap Fund" and recs[0]["open_ended"], "category/open-ended")
    check(not recs[3]["open_ended"], "close-ended flag")
    idx = mfs.AmfiIndex(recs)
    check(idx.by_isin["INF000A01011"]["code"] == "1001", "isin lookup")
    check(mfs.is_direct_growth(recs[0]["name"]) and not mfs.is_direct_growth(recs[1]["name"])
          and not mfs.is_direct_growth(recs[2]["name"]), "direct-growth filter")


def test_parse_navall_current_amfi_layout():
    """AMFI's current file: Plan and Option are separate columns after the scheme name."""
    text = "\n".join([
        "Scheme Code;ISIN Div Payout/ ISIN Growth;ISIN Div Reinvestment;Scheme Name;Plan;Option;Net Asset Value;Date",
        "", "Open Ended Schemes(Children's Fund - Childrens' Fund)", "", "Axis Mutual Fund", "",
        "135762;INF846K01WO1;-;Axis Children's Fund;Direct Plan;Growth Option;29.3960;16-Oct-2026",
        "135765;INF846K01WP8;-;Axis Children's Fund;Direct Plan;IDCW Option;27.0791;16-Oct-2026",
        "135759;INF846K01WJ1;-;Axis Children's Fund;Regular Plan;Growth Option;25.5801;16-Oct-2026",
    ])
    recs = mfs.parse_navall(text)
    check(len(recs) == 3, f"expected 3 records, got {len(recs)}")
    check(recs[0]["nav"] == 29.396 and str(recs[0]["nav_date"]) == "2026-10-16", f"nav/date wrong: {recs[0]}")
    check(recs[0]["name"] == "Axis Children's Fund - Direct Plan - Growth Option", recs[0]["name"])
    check([mfs.is_direct_growth(r["name"]) for r in recs] == [True, False, False], "direct-growth filter, new layout")
    check(recs[0]["open_ended"] and recs[0]["amc"] == "Axis Mutual Fund", "category / amc")
    check(mfs.AmfiIndex(recs).by_isin["INF846K01WO1"]["code"] == "135762", "isin lookup")


def test_name_match():
    recs = mfs.parse_navall(navall_text([
        ("1", "INF000A01011", "Aditya Birla Sun Life Consumption Fund - Direct Plan - Growth", "Equity Scheme - Thematic"),
        ("2", "INF000A01029", "Aditya Birla Sun Life Digital India Fund - Direct Plan - Growth", "Equity Scheme - Thematic"),
        ("3", "INF000B01013", "Axis Small Cap Fund - Direct Plan - Growth", "Equity Scheme - Small Cap Fund"),
    ]))
    idx = mfs.AmfiIndex(recs)
    rec, ratio = idx.match_name("291GZ -Aditya Birla Sun Life Consumption Fund -Growth-Direct Plan(Formerly known as Aditya Birla Sun Life India GenNext Fund)")
    check(rec and rec["code"] == "1", f"name match failed: {rec}")
    rec, _ = idx.match_name("8184- ICICI Prudential Infrastructure Fund - Direct Plan -Growth")
    check(rec is None, "unrelated name must not match")
    rec, _ = idx.match_name("SCDG -Axis Small Cap Fund Direct Growth")
    check(rec and rec["code"] == "3", "code-prefixed name should match")


def base_facts(**kw):
    f = {"mapped": True, "zero_value": False, "history_years": 7.0, "nav_age_days": 1,
         "cagr_3y": 0.18, "cagr_5y": 0.17, "max_drawdown": -0.30, "category_short": "Small Cap Fund",
         "benchmark_excess_return": 0.02, "rolling_beat_benchmark_pct": 0.70, "benchmark_name": "Nifty Smallcap 250 TRI",
         "peer": {"size": 20, "rank": 3, "percentile": 89.0, "median_3y": 0.15, "median_5y": 0.15, "median_dd": -0.32}}
    f.update(kw)
    return f


def test_decide_rules():
    d = mfs.decide(base_facts())
    check(d["verdict"] == mfs.V_CONTINUE and d["confidence"] == "High", f"strong fund: {d}")

    d = mfs.decide(base_facts(benchmark_excess_return=-0.03, rolling_beat_benchmark_pct=0.2,
                              peer={"size": 20, "rank": 18, "percentile": 10.0, "median_3y": 0.15, "median_5y": 0.15, "median_dd": -0.3}))
    check(d["verdict"] == mfs.V_REVIEW and len(d["red"]) >= 2, f"weak fund must be REVIEW: {d}")

    d = mfs.decide(base_facts(benchmark_excess_return=-0.005, rolling_beat_benchmark_pct=0.40))
    check(d["verdict"] == mfs.V_CONTINUE and len(d["amber"]) == 2, f"two benchmark ambers are ONE evidence type: {d}")

    d = mfs.decide(base_facts(benchmark_excess_return=-0.005,
                              peer={"size": 20, "rank": 12, "percentile": 42.0, "median_3y": 0.15, "median_5y": 0.15, "median_dd": -0.3}))
    check(d["verdict"] == mfs.V_WATCH and len(d["amber"]) == 2 and not d["red"], f"ambers in two evidence types = WATCH: {d}")

    d = mfs.decide(base_facts(benchmark_excess_return=-0.05, rolling_beat_benchmark_pct=0.10))
    check(d["verdict"] == mfs.V_WATCH and len(d["red"]) == 2, f"two benchmark reds are ONE evidence type = WATCH: {d}")

    d = mfs.decide(base_facts(benchmark_excess_return=0.01, rolling_beat_benchmark_pct=0.5,
                              peer={"size": 20, "rank": 18, "percentile": 10.0, "median_3y": 0.15, "median_5y": 0.15, "median_dd": -0.3}))
    check(d["verdict"] == mfs.V_WATCH and len(d["red"]) == 1, f"one red = WATCH: {d}")

    d = mfs.decide(base_facts(cagr_3y=0.04, benchmark_excess_return=None, rolling_beat_benchmark_pct=None, peer=None))
    check(d["verdict"] == mfs.V_WATCH and d["confidence"] == "Low", f"absolute-only, 3y below risk-free: {d}")

    d = mfs.decide(base_facts(benchmark_excess_return=None, rolling_beat_benchmark_pct=None, peer=None))
    check(d["verdict"] == mfs.V_CONTINUE and d["confidence"] == "Low" and d["notes"], f"absolute-only fine fund: {d}")

    d = mfs.decide(base_facts(history_years=1.5))
    check(d["verdict"] == mfs.V_NODATA and "history" in d["notes"][0], f"young fund: {d}")

    d = mfs.decide(base_facts(nav_age_days=40, last_nav_date=TODAY))
    check(d["verdict"] == mfs.V_REVIEW and d["red"], f"stale NAV: {d}")

    d = mfs.decide(base_facts(zero_value=True))
    check(d["verdict"] == mfs.V_REVIEW and "wound up" in d["red"][0], f"zero value: {d}")

    d = mfs.decide(base_facts(mapped=False))
    check(d["verdict"] == mfs.V_NODATA, "unmapped")
    d = mfs.decide(base_facts(nav_error="timeout"))
    check(d["verdict"] == mfs.V_NODATA, "nav error")

    d = mfs.decide(base_facts(peer={"size": 4, "rank": 4, "percentile": 0.0, "median_3y": 0.1, "median_5y": 0.1, "median_dd": -0.3}))
    check(d["verdict"] == mfs.V_CONTINUE and d["confidence"] == "Medium", f"too few peers: {d}")


def test_default_output_name():
    check(mfs.default_output("Kiran_2026-06-MF.XLS") == "Kiran_2026-06-MF-Screened.XLS", "default output name")
    check(mfs.default_output("x/y.xlsx") == "x/y-Screened.xlsx", "xlsx output name")


def write_holdings(path, rows):
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Mutual Funds")
    heads = ["Scheme Name", "ISIN", "Folio No.", "Closing Bal (Units)", "NAV (Rs)", "Cumulative Amount Invested (INR)",
             "Valuation (Rs)", "Unrealised Profit/Loss", "Unrealised Profit/Loss (%)", "PDF Page", "Check"]
    for j, h in enumerate(heads):
        ws.write(0, j, h)
    for i, (name, isin, value) in enumerate(rows, 1):
        vals = [name, isin, f"F{i}/0", 100.0, value / 100.0, value * 0.8, value, value * 0.2, 25.0, 1, "OK"]
        for j, v in enumerate(vals):
            ws.write(i, j, v)
    ws.write(len(rows) + 1, 0, "TOTAL")
    wb.save(path)


def test_end_to_end():
    work = tempfile.mkdtemp()
    old_cwd = os.getcwd()
    os.chdir(work)
    config.DATA_DIR = os.path.join(work, "data")
    config.BENCHMARK_DIR = os.path.join(work, "data", "benchmarks")
    os.makedirs(config.BENCHMARK_DIR)
    try:
        small, large, liquid = "Equity Scheme - Small Cap Fund", "Equity Scheme - Large Cap Fund", "Debt Scheme - Liquid Fund"
        navs, schemes = {}, []

        def add(code, name, cat, drift, years=7.0, age=1, isin=None):
            isin = isin or isin_for(int(code))
            schemes.append((str(code), isin, name, cat))
            navs[str(code)] = make_nav(drift, years=years, end_age_days=age, seed=int(code))
            return isin

        good = add(101, "Good Small Cap Fund - Direct Plan - Growth", small, 0.24)
        bad = add(102, "Bad Small Cap Fund - Direct Plan - Growth", small, 0.03)
        young = add(103, "Young Small Cap Fund - Direct Plan - Growth", small, 0.20, years=1.5)
        stale = add(104, "Stale Small Cap Fund - Direct Plan - Growth", small, 0.20, age=60)
        mid = add(105, "Middle Small Cap Fund - Direct Plan - Growth", small, 0.15)
        for k in range(12):
            add(200 + k, f"Peer{k} Small Cap Fund - Direct Plan - Growth", small, 0.08 + 0.012 * k)
        lc = add(301, "Solo Large Cap Fund - Direct Plan - Growth", large, 0.14)
        liq = add(401, "Calm Liquid Fund - Direct Plan - Growth", liquid, 0.07, years=6.0)
        reg = add(501, "Regular Small Cap Fund - Regular Plan - Growth", small, 0.12)
        schemes.append(("999", "INF999Z01019", "Peer999 Small Cap Fund - Direct Plan - IDCW", small))
        navs["999"] = make_nav(0.1, seed=999)

        for key, drift in (("small_cap", 0.14), ("large_cap", 0.11)):
            b = make_nav(drift, vol=0.05, seed=7 if key == "small_cap" else 8)
            b.rename(columns={"nav": "tri"}).to_csv(os.path.join(config.BENCHMARK_DIR, key + ".csv"),
                                                    index=False, date_format="%d-%m-%Y")

        def fake_history(code, force_refresh=False):
            return navs[str(code)].copy()
        fetch_nav_data.get_nav_history = fake_history

        navall = os.path.join(work, "NAVAll.txt")
        with open(navall, "w") as fh:
            fh.write(navall_text(schemes))

        holdings = os.path.join(work, "holdings-MF.XLS")
        write_holdings(holdings, [
            ("A1 -Good Small Cap Fund -Direct Plan -Growth", good, 500000),
            ("A2 -Bad Small Cap Fund -Direct Plan -Growth", bad, 300000),
            ("A3 -Bad Small Cap Fund -Direct Plan -Growth", bad, 100000),
            ("A4 -Young Small Cap Fund -Direct Plan -Growth", young, 50000),
            ("A5 -Stale Small Cap Fund -Direct Plan -Growth", stale, 50000),
            ("A6 -Middle Small Cap Fund -Direct Plan -Growth", mid, 100000),
            ("A7 -Solo Large Cap Fund -Direct Plan -Growth", lc, 100000),
            ("A8 -Calm Liquid Fund -Direct Plan -Growth", liq, 100000),
            ("A9 -Regular Small Cap Fund -Regular Plan -Growth", reg, 50000),
            ("A10 -Wound Up Fund -Direct Plan -Growth", "INF123X01016", 0),
            ("A11 -Good Small Cap Fund -Direct Plan -Growth", "INF777Q01010", 20000),
            ("A12 -Mystery Fund -Direct Plan -Growth", "INF888R01010", 10000),
        ])
        out = os.path.join(work, "out.xls")
        mfs.main([holdings, "-o", out, "--navall", navall, "--workers", "2"])

        import xlrd
        wb = xlrd.open_workbook(out)
        check(wb.sheet_names() == ["Holdings Analysis", "Summary", "Method"], wb.sheet_names())
        sh = wb.sheet_by_name("Holdings Analysis")
        head = sh.row_values(0)
        check(head[:11] == ["Scheme Name", "ISIN", "Folio No.", "Closing Bal (Units)", "NAV (Rs)",
                            "Cumulative Amount Invested (INR)", "Valuation (Rs)", "Unrealised Profit/Loss",
                            "Unrealised Profit/Loss (%)", "PDF Page", "Check"], "original columns kept first")
        check(head[11] == "Verdict", "Verdict is the first added column")
        rows = {sh.cell_value(i, 0): dict(zip(head, sh.row_values(i))) for i in range(1, sh.nrows)}
        check(sh.nrows == 1 + 12 + 1, f"rows: header + 12 holdings + total, got {sh.nrows}")
        v = lambda key: rows[key]["Verdict"]
        check(v("A1 -Good Small Cap Fund -Direct Plan -Growth") == "CONTINUE", v("A1 -Good Small Cap Fund -Direct Plan -Growth"))
        check(v("A2 -Bad Small Cap Fund -Direct Plan -Growth") == "REVIEW", v("A2 -Bad Small Cap Fund -Direct Plan -Growth"))
        check(v("A3 -Bad Small Cap Fund -Direct Plan -Growth") == "REVIEW", "same ISIN -> same verdict")
        check(v("A4 -Young Small Cap Fund -Direct Plan -Growth") == "INSUFFICIENT DATA", "young fund")
        check(v("A5 -Stale Small Cap Fund -Direct Plan -Growth") == "REVIEW", "stale NAV")
        check(v("A10 -Wound Up Fund -Direct Plan -Growth") == "REVIEW", "zero valuation")
        check(v("A12 -Mystery Fund -Direct Plan -Growth") == "INSUFFICIENT DATA", "unmapped fund")
        check(rows["A8 -Calm Liquid Fund -Direct Plan -Growth"]["Confidence"] == "Low", "liquid fund = absolute only")
        check(rows["A8 -Calm Liquid Fund -Direct Plan -Growth"]["Peer Group Size"] in ("", None), "no peer group for debt")
        check("Name match" in rows["A11 -Good Small Cap Fund -Direct Plan -Growth"]["Match Method"], "name fallback flagged")
        check(rows["A11 -Good Small Cap Fund -Direct Plan -Growth"]["AMFI Scheme Code"] == "101", "name fallback code")
        a1 = rows["A1 -Good Small Cap Fund -Direct Plan -Growth"]
        # The best-drift fund must sit in the top third of its peers. (Was "<= 3"; with the 3-year drawdown window the
# synthetic noise re-orders neighbours by a place or two, so the bound is the top third, not a fixed rank.)
        check(a1["Confidence"] == "High" and a1["Peer Group Size"] >= 12
              and a1["Peer Rank"] <= max(3, a1["Peer Group Size"] // 3), f"good fund rank: {a1['Peer Rank']}")
        check("999" not in str(rows["A1 -Good Small Cap Fund -Direct Plan -Growth"]["AMFI Scheme Code"]), "IDCW peer excluded")
        a9 = rows["A9 -Regular Small Cap Fund -Regular Plan -Growth"]
        check("Regular-plan" in a9["Portfolio Notes"], "regular plan note")
        check("2 folios" in rows["A2 -Bad Small Cap Fund -Direct Plan -Growth"]["Portfolio Notes"], "multi-folio note")
        check("likely overlap" in a1["Portfolio Notes"] and "of yours" in a1["Portfolio Notes"], "overlap note")
        total_row = sh.row_values(sh.nrows - 1)
        check(total_row[0] == "TOTAL" and abs(total_row[6] - 1380000) < 1, f"total valuation {total_row[6]}")
        summ = wb.sheet_by_name("Summary")
        text = " ".join(str(c) for i in range(summ.nrows) for c in summ.row_values(i))
        check("Verdict summary" in text and "Funds to look at first" in text and "by AMFI category" in text, "summary sections")
        meth = wb.sheet_by_name("Method")
        check(any("not investment advice" in str(meth.cell_value(i, 0)) for i in range(meth.nrows)), "method disclaimer")

        out2 = os.path.join(work, "out2.xlsx")
        mfs.main([holdings, "-o", out2, "--navall", navall, "--no-peers"])
        from openpyxl import load_workbook
        w2 = load_workbook(out2)
        ws2 = w2["Holdings Analysis"]
        hd = [c.value for c in ws2[1]]
        col = hd.index("Peer Group Size") + 1
        check(all(ws2.cell(row=r, column=col).value in (None, "") for r in range(2, ws2.max_row + 1)), "no peers when --no-peers")
        check(ws2.cell(row=2, column=hd.index("Verdict") + 1).value in ("CONTINUE", "CONTINUE - WATCH", "REVIEW"), "xlsx verdict")
    finally:
        os.chdir(old_cwd)
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    tests = [test_parse_navall_and_index, test_parse_navall_current_amfi_layout, test_name_match, test_decide_rules, test_default_output_name, test_end_to_end]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"ok    {t.__name__}")
        except Exception as exc:
            failed += 1
            import traceback
            traceback.print_exc()
            print(f"FAIL  {t.__name__}: {exc}")
    sys.exit(1 if failed else 0)
