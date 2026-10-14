"""Offline tests for parameter 8: portfolio_quality.py, holdings_fetch.py, amc_registry.py
and their wiring into shortlist.py. Uses synthetic SEBI-style workbooks; no network."""
import os
import sys
import tempfile
import types
import zipfile
from datetime import datetime, timedelta

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import config
import amc_registry
import holdings_fetch
import portfolio_quality as pq

passed = 0


def check(name, cond):
    global passed
    assert cond, f"FAIL: {name}"
    passed += 1
    print(f"ok  {name}")


TODAY = datetime.now()
FRESH = (TODAY - timedelta(days=20)).strftime("%d-%b-%Y")
STALE = (TODAY - timedelta(days=200)).strftime("%d-%b-%Y")
INDUSTRIES = ["Banks", "IT - Software", "Auto Components", "Chemicals", "Healthcare Services",
              "Consumer Durables", "Capital Goods", "Finance"]


def equities(n, top10_each=3.0, total=95.0):
    """n equity holdings: ten at top10_each, the rest share what is left of `total`."""
    rest = (total - 10 * top10_each) / (n - 10) if n > 10 else 0
    out = []
    for i in range(n):
        out.append((f"Company {i}", f"INE{i:03d}A01{i:03d}", INDUSTRIES[i % len(INDUSTRIES)],
                    top10_each if i < 10 else rest))
    return out


def write_sheet(ws, title, as_of, eq, cash=(("TREPS", 3.0), ("Net Current Assets", 2.0)), scale=1.0,
                futures=True):
    ws.append(["Axis Mutual Fund"])
    ws.append([title])
    ws.append([f"Monthly Portfolio Statement as on {as_of}"])
    ws.append([])
    ws.append(["Name of the Instrument", "ISIN", "Industry / Rating", "Quantity",
               "Market/Fair Value (Rs. in Lakhs)", "% to NAV"])
    ws.append(["Equity & Equity related"])
    for name, isin, ind, pct in eq:
        ws.append([name, isin, ind, 1000, 50.0, pct * scale])
    if futures:
        ws.append(["Index Futures Dec 2026", "", "", 10, 5.0, 2.0 * scale])
    ws.append(["Sub Total", None, None, None, None, sum(e[3] for e in eq) * scale])
    ws.append(["Money Market Instruments"])
    for name, pct in cash:
        ws.append([name, "", "", None, 1.0, pct * scale])
    ws.append(["Grand Total", None, None, None, None, 100.0 * scale])
    ws.append(["Company Zzz After Total", "INE999A01999", "Banks", 1, 1, 50.0])  # must be ignored


def expected_score(eq, cash_pct, mcap_share=None):
    """Independent re-computation of the rubric from config.PQ_CRITERIA."""
    crit = config.PQ_CRITERIA
    w = np.array([e[3] for e in eq])
    ind = pd.Series([e[2] for e in eq])
    inputs = {"top10_pct": float(np.sort(w)[::-1][:10].sum()), "n_holdings": float(len(eq)),
              "cash_pct": cash_pct,
              "top_sector_pct": float(pd.Series(w).groupby(ind).sum().max())}
    if mcap_share is not None:
        inputs["largemid_pct"] = mcap_share
    num = den = 0.0
    for k, x in inputs.items():
        c = crit[k]
        s = min(1.0, max(0.0, (x - c["worst"]) / (c["best"] - c["worst"]))) * 100
        num += s * c["weight"]
        den += c["weight"]
    return num / den, inputs


EQ_A = equities(40)
EQ_C = equities(6, top10_each=10.0, total=60.0)  # note: only 6 holdings


with tempfile.TemporaryDirectory() as td:
    config.HOLDINGS_DIR = td
    config.AMFI_MARKETCAP_CSV = os.path.join(td, "no_such.csv")

    # ---- build workbooks ---------------------------------------------------
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_sheet(wb.create_sheet("AXSC"), "Axis Small Cap Fund (An open ended equity scheme predominantly "
                "investing in small cap stocks)", FRESH, EQ_A)
    write_sheet(wb.create_sheet("FRAC"), "Fraction Style Fund (An open ended equity scheme)", FRESH, EQ_A, scale=0.01)
    write_sheet(wb.create_sheet("FEW"), "Few Stocks Fund (An open ended equity scheme)", FRESH, EQ_C,
                cash=(("TREPS", 30.0), ("Net Current Assets", 10.0)))
    write_sheet(wb.create_sheet("BADSUM"), "Bad Sum Fund (An open ended equity scheme)", FRESH, EQ_A,
                cash=(("TREPS", 3.0),), scale=0.8)
    write_sheet(wb.create_sheet("OLD"), "Old Statement Fund (An open ended equity scheme)", STALE, EQ_A)
    ws = wb.create_sheet("Index")  # a non-portfolio sheet must be ignored, not crash
    ws.append(["Scheme", "Sheet"])
    ws.append(["Axis Small Cap Fund", "AXSC"])
    wb.save(os.path.join(td, "axis_month.xlsx"))

    # a zip holding one more workbook, with a path-traversal member name
    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    write_sheet(wb2.create_sheet("ZIPD"), "Zipped Fund (An open ended equity scheme)", FRESH, EQ_A)
    inner = os.path.join(td, "_inner.xlsx")
    wb2.save(inner)
    with zipfile.ZipFile(os.path.join(td, "dsp_month.zip"), "w") as z:
        z.write(inner, "../../evil_zipped.xlsx")
    os.remove(inner)

    idx = pq.PortfolioIndex()
    check("workbooks read (xlsx + one inside zip)", idx.files_read == 2)
    check("6 portfolio sheets parsed, Index sheet ignored", len(idx.entries) == 6)
    check("zip member stayed inside extraction dir (no path traversal)",
          not os.path.exists(os.path.join(os.path.dirname(td), "evil_zipped.xlsx"))
          and not os.path.exists(os.path.join(td, "evil_zipped.xlsx")))

    # ---- scoring -----------------------------------------------------------
    exp, inputs = expected_score(EQ_A, 5.0)
    a = pq.fields_for_fund(idx, "Axis Small Cap Fund - Direct Plan - Growth")
    check("matched by exact normalised name", a["pq_source"] == "axis_month.xlsx")
    check("score equals independent rubric calculation", abs(a["portfolio_quality"] - exp) < 1e-6)
    check("top10 measured (rows after Grand Total ignored)", abs(a["pq_top10_pct"] - inputs["top10_pct"]) < 1e-9)
    check("holdings count excludes futures/cash", a["pq_holdings"] == 40)
    check("cash = TREPS + net current assets", abs(a["pq_cash_pct"] - 5.0) < 1e-9)
    check("top sector matches", abs(a["pq_top_sector_pct"] - inputs["top_sector_pct"]) < 1e-9)
    check("as-of parsed", a["pq_as_of"] == pd.to_datetime(FRESH, dayfirst=True).strftime("%Y-%m-%d"))
    check("missing largemid is None, coverage 75%", a["pq_largemid_pct"] is None and abs(a["pq_coverage"] - 0.75) < 1e-9)

    f = pq.fields_for_fund(idx, "Fraction Style Fund - Direct Plan - Growth")
    check("fraction-style weights (0.05) give the same score", abs(f["portfolio_quality"] - exp) < 1e-6)

    few = pq.fields_for_fund(idx, "Few Stocks Fund - Direct Plan - Growth")
    check("<10 equity holdings -> None with reason", few["portfolio_quality"] is None and "equity holdings" in few["pq_note"])

    bad = pq.fields_for_fund(idx, "Bad Sum Fund - Direct Plan - Growth")
    check("weights not ~100% -> None, never a guess", bad["portfolio_quality"] is None and "sum" in bad["pq_note"])

    old = pq.fields_for_fund(idx, "Old Statement Fund - Direct Plan - Growth")
    check("stale statement -> None with reason", old["portfolio_quality"] is None and "stale" in old["pq_note"])

    z = pq.fields_for_fund(idx, "Zipped Fund - Direct Plan - Growth")
    check("fund inside a zip is scored", z["portfolio_quality"] is not None)

    none = pq.fields_for_fund(idx, "Completely Different Fund - Direct Plan - Growth")
    check("unmatched scheme -> None + note", none["portfolio_quality"] is None and "no holdings file" in none["pq_note"])
    fuzzy = pq.fields_for_fund(idx, "Axis Smallcap Fund - Direct Plan - Growth")
    check("near-identical spelling matches (Smallcap vs Small Cap)", fuzzy["portfolio_quality"] is not None)
    other = pq.fields_for_fund(idx, "Axis Mid Cap Fund - Direct Plan - Growth")
    check("different scheme from same AMC does NOT match", other["portfolio_quality"] is None)
    check("index None -> full key set", set(pq.fields_for_fund(None, "x")) == set(pq.EMPTY_FIELDS))

    # ---- market-cap sub-criterion -----------------------------------------
    isins = [e[1] for e in EQ_A]
    cats = ["large"] * 20 + ["mid"] * 10 + ["small"] * 10
    mc = dict(zip(isins, cats))
    q = pq.compute_quality(idx.lookup("Axis Small Cap Fund")["holdings"], mc)
    w = np.array([e[3] for e in EQ_A])
    share = w[:30].sum() / w.sum() * 100
    exp_mc, _ = expected_score(EQ_A, 5.0, mcap_share=share)
    check("largemid_pct measured from AMFI map", abs(q["inputs"]["largemid_pct"] - share) < 1e-9)
    check("score with market-cap criterion matches rubric", abs(q["score"] - exp_mc) < 1e-6 and q["coverage"] == 1.0)
    partial = dict(list(mc.items())[:20])  # only half the equity weight found -> below PQ_MIN_MCAP_MATCH
    q2 = pq.compute_quality(idx.lookup("Axis Small Cap Fund")["holdings"], partial)
    check("poor AMFI-list coverage -> criterion skipped, not guessed", "largemid_pct" not in q2["inputs"])

    # ---- coverage floor ----------------------------------------------------
    saved = config.PQ_MIN_COVERAGE
    config.PQ_MIN_COVERAGE = 0.9
    q3 = pq.compute_quality(idx.lookup("Axis Small Cap Fund")["holdings"], None)
    config.PQ_MIN_COVERAGE = saved
    check("coverage below floor -> score None", q3["score"] is None and "measurable" in q3["reason"])

    # ---- name / date helpers ----------------------------------------------
    check("norm strips plan words and parentheses",
          pq.norm_scheme_name("Axis Small Cap Fund (An open ended scheme) - Direct Plan - Growth") == "axis small cap")
    check("classify: equity ISIN", pq.classify_row("X Ltd", "INE002A01018") == "equity")
    check("classify: debenture ISIN is not equity", pq.classify_row("X NCD", "INE002A07123") == "other")
    check("classify: TREPS is cash", pq.classify_row("TREPS - 02 Oct 2026", "") == "cash")
    check("classify: futures is derivative", pq.classify_row("HDFC Bank Ltd Oct Futures", "") == "derivative")

    # ---- shortlist wiring (stubbed data layer) -----------------------------
    fake_fu = types.ModuleType("fund_universe"); fake_fu.supported_categories = lambda: []
    fake_fetch = types.ModuleType("fetch_nav_data")
    dates = pd.bdate_range("2019-01-01", TODAY.strftime("%Y-%m-%d"))
    nav = pd.DataFrame({"date": dates, "nav": 100 * np.exp(np.arange(len(dates)) * 0.0004)})
    fake_fetch.get_nav_history = lambda code, force_refresh=False: nav
    names = {"1": "Axis Small Cap Fund - Direct Plan - Growth", "2": "Unlisted Small Cap Fund - Direct Plan - Growth",
             "3": "Few Stocks Fund - Direct Plan - Growth"}
    fake_fetch.get_scheme_meta = lambda code: {"scheme_name": names[str(code)], "fund_house": "Axis Mutual Fund",
                                               "scheme_category": "Equity Scheme - Small Cap Fund"}
    fake_metrics = types.ModuleType("metrics")
    rng = np.random.default_rng(3)
    fake_metrics.compute_all_metrics = lambda *a, **k: {
        "cagr_1y": rng.uniform(.1, .3), "cagr_3y": rng.uniform(.1, .3), "cagr_5y": rng.uniform(.1, .3),
        "volatility": .2, "max_drawdown": -rng.uniform(.2, .4), "risk_return_ratio": rng.uniform(.5, 1.5),
        "downside_deviation": .1, "sortino_ratio": rng.uniform(.5, 2), "rolling_avg": .15, "rolling_min": .02,
        "rolling_positive_pct": rng.uniform(.6, 1), "rolling_std": .05, "sip_xirr": rng.uniform(.1, .3),
        "history_years": 7.0}
    sys.modules.update(fund_universe=fake_fu, fetch_nav_data=fake_fetch, metrics=fake_metrics)
    config.BENCHMARK_ENABLED = False
    import shortlist

    flist = pd.DataFrame({"scheme_code": ["1", "2", "3"]})
    table = shortlist.build_metrics_table(flist)
    got = dict(zip(table["scheme_code"], table["portfolio_quality"]))
    check("pipeline: matched fund scored, others missing",
          pd.notna(got["1"]) and pd.isna(got["2"]) and pd.isna(got["3"]))
    check("pipeline: pq_note explains misses",
          "no holdings file" in table.loc[table.scheme_code == "2", "pq_note"].iloc[0]
          and "equity holdings" in table.loc[table.scheme_code == "3", "pq_note"].iloc[0])
    ranked = shortlist.score_and_rank(table)
    check("ranking works with partly-missing portfolio_quality", ranked["composite_score"].notna().all())
    disp = shortlist.format_for_display(ranked)
    check("display shows N/A for missing score", (disp["portfolio_quality"] == "N/A").sum() == 2)

    # an empty holdings folder must change nothing (all missing, still ranks)
    empty = tempfile.mkdtemp()
    config.HOLDINGS_DIR = empty
    t2 = shortlist.build_metrics_table(flist)
    check("no holdings files -> every score missing, no crash", t2["portfolio_quality"].isna().all())
    check("no holdings files -> ranking still works", shortlist.score_and_rank(t2)["composite_score"].notna().all())

# ---- holdings_fetch helpers --------------------------------------------------
html_page = '''<a href="/files/NIMF-MONTHLY-PORTFOLIO-31-Aug-26.xls">a</a>
<a href='https://x.com/f/NIMF-MONTHLY-PORTFOLIO-30-Sep-26.xls?v=1'>b</a><a href="/files/factsheet.pdf">c</a>
<a href="/files/NIMF-MONTHLY-PORTFOLIO-31-Aug-26.xls">dup</a>'''
links = holdings_fetch.extract_links(html_page, "https://mf.example.com/page")
check("extract_links absolute + de-duplicated", len(links) == 3 and links[0].startswith("https://mf.example.com/files/"))
check("date from 'dd-Mon-yy'", holdings_fetch.date_from_filename(links[1]) == datetime(2026, 9, 30))
check("date from dd-mm-yyyy",
      holdings_fetch.date_from_filename("https://a/Monthly_Portfolio_30-09-2026.xls") == datetime(2026, 9, 30))
picked = holdings_fetch.pick_nippon(links)
check("nippon picker ignores non-portfolio links", len(picked) == 2)
check("newest() returns the latest month", "30-Sep-26" in holdings_fetch.newest(picked, 1)[0])
check("baroda picker", len(holdings_fetch.pick_baroda(["https://b/Monthly_Portfolio_30-09-2026.xls", "https://b/x.xls"])) == 1)
check("dsp picker (.zip with month+portfolio)",
      len(holdings_fetch.pick_dsp(["https://d/Month-end-Portfolio-Sep.zip", "https://d/other.zip"])) == 1)
check("HTML error page is rejected as non-Excel", not holdings_fetch.looks_like_office_or_zip(b"<html>blocked</html>" * 200))
check("xlsx/zip/xls magic bytes accepted",
      holdings_fetch.looks_like_office_or_zip(b"PK\x03\x04") and holdings_fetch.looks_like_office_or_zip(b"\xd0\xcf\x11\xe0"))


class FakeResp:
    def __init__(self, content): self.content = content; self.text = ""
    def raise_for_status(self): pass


with tempfile.TemporaryDirectory() as dd:
    orig = holdings_fetch.requests.get
    holdings_fetch.requests.get = lambda *a, **k: FakeResp(b"<html>Access denied</html>" * 300)
    p, msg = holdings_fetch.download("https://x/a.xls", dd)
    check("download refuses an HTML page named .xls", p is None and not os.listdir(dd))
    holdings_fetch.requests.get = lambda *a, **k: FakeResp(b"PK\x03\x04" + b"0" * 5000)
    p, msg = holdings_fetch.download("https://x/b.xlsx", dd)
    check("download saves a real zip/xlsx payload", p and os.path.exists(p) and msg == "downloaded")
    p2, msg2 = holdings_fetch.download("https://x/b.xlsx", dd)
    check("second download is skipped", msg2 == "already downloaded")
    holdings_fetch.requests.get = orig

# ---- registry ------------------------------------------------------------------
keys = [a["key"] for a in amc_registry.AMCS]
check("registry covers all 32 requested AMCs, unique keys", len(keys) == 32 and len(set(keys)) == 32)
check("every automated AMC has a link picker", all(k in holdings_fetch.PICKERS
      for k in keys if amc_registry.BY_KEY[k]["mode"] == "auto"))
check("only known modes", {a["mode"] for a in amc_registry.AMCS} <=
      {"auto", "manual", "manual_unconfirmed", "gated", "not_found"})
for house, want in [("Kotak Mahindra Mutual Fund", "kotak"), ("Mahindra Manulife Mutual Fund", "mahindra_manulife"),
                    ("Axis Mutual Fund", "axis"), ("Bank of India Mutual Fund", "bank_of_india"),
                    ("quant Mutual Fund", "quant"), ("Nippon India Mutual Fund", "nippon_india"),
                    ("SBI Mutual Fund", "sbi"), ("Baroda BNP Paribas Mutual Fund", "baroda_bnp_paribas")]:
    got = amc_registry.amc_for_fund_house(house)
    check(f"fund-house match: {house}", got is not None and got["key"] == want)
check("fund house not in registry -> None (e.g. IDBI)", amc_registry.amc_for_fund_house("IDBI Mutual Fund") is None)

print(f"\nAll {passed} checks passed")
