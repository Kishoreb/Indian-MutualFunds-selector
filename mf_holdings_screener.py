#!/usr/bin/env python3
"""
mf_holdings_screener.py
=======================
Screens every mutual fund you HOLD and says whether to keep holding it.

Input : the mutual fund holdings .xls produced by cdsl_holdings_to_xls.py
        (e.g. Kiran_2026-06-MF.XLS).  .xlsx also works.
Output: the same rows plus extra analysis columns, written to a new .xls
        (<input name>-Screened.XLS by default), with a Summary sheet and a
        Method sheet that lists every rule and threshold used.

For each holding it:
  1. maps the ISIN to an AMFI scheme code and category (AMFI's NAVAll.txt);
  2. downloads the NAV history (mfapi.in, same source as shortlist.py) and computes the
     toolkit's usual metrics (metrics.py): 1/3/5y CAGR, volatility, max drawdown, Sortino,
     rolling 3y consistency, SIP XIRR;
  3. compares it with its category benchmark (benchmark.py): 3y excess return and the share
     of rolling 3y windows in which it beat the benchmark;
  4. ranks it against every Direct-Growth peer in the same AMFI category using shortlist.py's
     scoring engine and the weights in weights.ini (metrics it cannot compute here are left
     out and the other weights rescaled; a fund with too little data is not ranked);
  5. turns those into red / amber / green flags and a verdict:
         CONTINUE | CONTINUE - WATCH | REVIEW | INSUFFICIENT DATA
  6. adds portfolio-level notes: how much of your MF money a fund is, and how many of your
     funds sit in the same category (overlap), and which of them ranks best.

It is a rules-based screen of PAST data, not investment advice.  It ignores exit load and
tax on selling, your goals and your holding period - check those before acting.  Thresholds
live in REVIEW_RULES below.

Usage
-----
    python mf_holdings_screener.py Kiran_2026-06-MF.XLS
    python mf_holdings_screener.py Kiran_2026-06-MF.XLS -o my_review.xls
    python mf_holdings_screener.py Kiran_2026-06-MF.XLS --no-peers      # faster; no peer ranking
    python mf_holdings_screener.py Kiran_2026-06-MF.XLS --navall NAVAll.txt   # offline AMFI file
    python mf_holdings_screener.py Kiran_2026-06-MF.XLS --refresh

Run it in the folder that holds the toolkit (config.py, metrics.py, benchmark.py,
fetch_nav_data.py, shortlist.py ...) on a computer with normal internet access.
"""

import argparse
import difflib
import math
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta

import pandas as pd
import requests

import benchmark
import config
import fetch_nav_data
from metrics import compute_all_metrics

# --------------------------------------------------------------------------
# Rules.  Every threshold the verdict uses is here (fractions: 0.02 = 2 percentage points).
# --------------------------------------------------------------------------
REVIEW_RULES = {
    "min_history_years": 3.0,          # less NAV history than this -> INSUFFICIENT DATA (too young to judge)
    "min_peer_history_years": 3.0,     # peers need this much history to be ranked
    "min_peer_group_size": 8,          # fewer ranked peers than this -> peer evidence not used
    "peer_red_percentile": 25,         # peer percentile (100 = best) at or below this -> red flag
    "peer_amber_percentile": 50,       # below this (but above red) -> amber flag
    "peer_green_percentile": 50,       # at or above this -> positive
    "bench_excess_amber": 0.0,         # 3y excess return vs benchmark below this -> amber
    "bench_excess_red": -0.02,         # ... at or below this -> red (2 percentage points a year behind)
    "rolling_beat_amber": 0.45,        # beat the benchmark in fewer than this share of rolling windows -> amber
    "rolling_beat_red": 0.30,          # ... fewer than this -> red
    "rolling_beat_green": 0.60,        # at or above this -> positive
    "drawdown_worse_than_peers": 0.10, # max drawdown this much deeper than the peer median -> amber
    "stale_nav_days": 10,              # last NAV older than this -> scheme may be closed/merged -> REVIEW
    "max_single_fund_weight": 0.15,    # one fund above this share of your MF portfolio -> note
    "overlap_funds": 3,                # this many of your funds in one category -> overlap note
}
# AMFI categories that are fair peer groups (matched on the lower-cased category text).
# Index funds, fund-of-funds, sectoral/thematic and debt categories mix very different
# mandates, so they are judged on absolute numbers only.
PEER_CATEGORY_KEYWORDS = [
    "large cap fund", "mid cap fund", "small cap fund", "large & mid cap fund",
    "multi cap fund", "flexi cap fund", "elss", "focused fund", "value fund", "contra fund",
    "dividend yield fund", "aggressive hybrid", "balanced hybrid", "conservative hybrid",
    "dynamic asset allocation", "multi asset allocation", "equity savings", "arbitrage fund",
]
NAVALL_URLS = ["https://portal.amfiindia.com/spages/NAVAll.txt",
               "https://www.amfiindia.com/spages/NAVAll.txt"]
NAME_MATCH_MIN = 0.90      # name fallback: minimum similarity
NAME_MATCH_MARGIN = 0.02   # ... and lead over the runner-up

V_CONTINUE, V_WATCH, V_REVIEW, V_NODATA = "CONTINUE", "CONTINUE - WATCH", "REVIEW", "INSUFFICIENT DATA"


# --------------------------------------------------------------------------
# AMFI NAVAll.txt: ISIN -> scheme code, name, category
# --------------------------------------------------------------------------
_HEAD_RE = re.compile(r"^(Open Ended|Close Ended|Interval Fund)[^()]*\((.*)\)\s*$", re.I)
_ISIN_RE = re.compile(r"^IN[A-Z0-9]{9}\d$")
_NON_GROWTH = ("idcw", "dividend", "bonus", "payout", "reinvest", "regular", "institutional")


def _navall_columns(header_line):
    """Column positions from AMFI's header row.  AMFI now sends the plan and option in their own
    columns (Scheme Name;Plan;Option;Net Asset Value;Date); older files put them in the name
    (Scheme Name;Net Asset Value;Date).  Reading the header copes with both."""
    cols = [c.strip().lower() for c in header_line.split(";")]

    def find(*keys):
        for i, c in enumerate(cols):
            if any(k in c for k in keys):
                return i
        return None
    return {"name": find("scheme name"), "plan": find("plan") if "plan" in cols else None,
            "option": cols.index("option") if "option" in cols else None,
            "nav": find("net asset value", "nav"), "date": find("date")}


def parse_navall(text):
    """AMFI NAVAll.txt -> list of dicts: code, isins, name, category, open_ended, amc, nav, nav_date.
    When the file has separate Plan / Option columns they are appended to the name
    ('Axis Children's Fund - Direct Plan - Growth Option') so is_direct_growth() still works."""
    recs, category, open_ended, amc = [], "", False, ""
    cols = {"name": 3, "plan": None, "option": None, "nav": 4, "date": 5}      # old layout
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if ";" in line:
            f = [x.strip() for x in line.split(";")]
            if f[0].lower().startswith("scheme code"):
                cols = _navall_columns(line)
                if cols["name"] is None or cols["nav"] is None:
                    cols = {"name": 3, "plan": None, "option": None, "nav": 4, "date": 5}
                continue
            if len(f) < 5 or not f[0].isdigit():
                continue

            def at(key):
                i = cols[key]
                return f[i] if i is not None and i < len(f) else ""
            name = at("name")
            extras = [x for x in (at("plan"), at("option")) if x]
            if extras:
                name = " - ".join([name] + extras)
            nav, nav_date = None, None
            try:
                nav = float(at("nav"))
            except ValueError:
                pass
            try:
                nav_date = datetime.strptime(at("date"), "%d-%b-%Y").date()
            except ValueError:
                pass
            recs.append({"code": f[0], "isins": [x for x in (f[1], f[2]) if _ISIN_RE.match(x)],
                         "name": name, "category": category, "open_ended": open_ended,
                         "amc": amc, "nav": nav, "nav_date": nav_date})
            continue
        m = _HEAD_RE.match(line)
        if m:
            category, open_ended = m.group(2).strip(), line.lower().startswith("open ended")
            continue
        amc = line
    return recs


def is_direct_growth(name):
    n = name.lower()
    return "direct" in n and "growth" in n and not any(w in n for w in _NON_GROWTH)


def norm_name(s, strip_code=False):
    """Name normalised for fuzzy matching.  strip_code removes the scheme code the CAS prints
    in front of the name ('8184- ICICI Prudential ...')."""
    s = re.sub(r"\([^)]*\)", " ", s)
    if strip_code:
        s = re.sub(r"^\s*[A-Z0-9]{1,6}(?:\s[A-Z0-9]{1,5})?\s*-\s*(?=\S)", "", s)
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    drop = {"direct", "plan", "growth", "option", "regular", "scheme", "the"}
    return " ".join(t for t in s.split() if t not in drop)


class AmfiIndex:
    def __init__(self, recs):
        self.records = recs
        self.by_isin = {}
        for r in recs:
            for i in r["isins"]:
                self.by_isin.setdefault(i, r)
        self.by_code = {r["code"]: r for r in recs}
        self._dg = None

    def _direct_growth(self):
        if self._dg is None:
            self._dg = [(norm_name(r["name"]), r) for r in self.records if is_direct_growth(r["name"])]
        return self._dg

    def match_name(self, cas_name):
        """Best Direct-Growth scheme by name, or (None, 0) if none is clearly best."""
        target = norm_name(cas_name, strip_code=True)
        if not target:
            return None, 0.0
        scored = sorted(((difflib.SequenceMatcher(None, target, n).ratio(), r["code"], r)
                         for n, r in self._direct_growth()), key=lambda x: -x[0])
        if not scored or scored[0][0] < NAME_MATCH_MIN:
            return None, 0.0
        if len(scored) > 1 and scored[0][0] - scored[1][0] < NAME_MATCH_MARGIN:
            return None, 0.0
        return scored[0][2], scored[0][0]


def load_navall(path=None, force=False, log=print):
    """Text of AMFI's NAVAll.txt: from `path`, a fresh cache, or a download."""
    if path:
        with open(path, encoding="utf-8", errors="replace") as fh:
            return fh.read()
    os.makedirs(config.DATA_DIR, exist_ok=True)
    cache = os.path.join(config.DATA_DIR, "NAVAll.txt")
    if not force and os.path.exists(cache):
        age = datetime.now() - datetime.fromtimestamp(os.path.getmtime(cache))
        if age < timedelta(days=config.CACHE_MAX_AGE_DAYS):
            with open(cache, encoding="utf-8", errors="replace") as fh:
                return fh.read()
    last = None
    for url in NAVALL_URLS:
        try:
            r = requests.get(url, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
            r.raise_for_status()
            text = r.content.decode("utf-8", errors="replace")   # r.text may guess latin-1 and garble apostrophes
            if ";" not in text:
                raise ValueError("unexpected content")
            with open(cache, "w", encoding="utf-8") as fh:
                fh.write(text)
            return text
        except Exception as exc:                      # try the next URL
            last = exc
            log(f"  could not download {url}: {exc}")
    if os.path.exists(cache):
        log("  using the older cached NAVAll.txt")
        with open(cache, encoding="utf-8", errors="replace") as fh:
            return fh.read()
    sys.exit(f"Could not get AMFI's NAVAll.txt ({last}).\n"
             f"Download it in a browser from {NAVALL_URLS[0]} and pass it with --navall.")


# --------------------------------------------------------------------------
# Reading the holdings file
# --------------------------------------------------------------------------
def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_holdings(path):
    """-> (headers, rows).  rows are dicts keyed by header; the TOTAL row and blanks are skipped."""
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        grid = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        import xlrd
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        grid = [sh.row_values(i) for i in range(sh.nrows)]
    hdr_i = next((i for i, r in enumerate(grid) if any(str(c).strip() == "ISIN" for c in r)), None)
    if hdr_i is None:
        sys.exit("No header row with an 'ISIN' column was found - is this the mutual fund holdings file?")
    headers = [str(c).strip() for c in grid[hdr_i]]
    need = ["Scheme Name", "ISIN"]
    if any(n not in headers for n in need):
        sys.exit(f"The holdings file needs the columns {need}; found {headers}")
    rows = []
    for r in grid[hdr_i + 1:]:
        d = {h: (c if c != "" else None) for h, c in zip(headers, r)}
        name = str(d.get("Scheme Name") or "").strip()
        if not name or name.upper() == "TOTAL":
            continue
        d["Scheme Name"], d["ISIN"] = name, str(d.get("ISIN") or "").strip().upper()
        rows.append(d)
    return headers, rows


# --------------------------------------------------------------------------
# Analysis of one fund
# --------------------------------------------------------------------------
def peer_eligible(category):
    c = (category or "").lower()
    return any(k in c for k in PEER_CATEGORY_KEYWORDS)


def analyse_nav(nav_df, category, provider, today):
    nav_df = nav_df.sort_values("date").reset_index(drop=True)
    m = compute_all_metrics(nav_df, config.RISK_FREE_RATE, sip_years=config.SIP_YEARS,
                            sip_monthly_amount=config.SIP_MONTHLY_AMOUNT,
                            rolling_window_years=config.ROLLING_WINDOW_YEARS,
                            align_ratio_window=config.RATIO_ALIGN_WINDOW)
    m.update(benchmark.metrics_for_fund(nav_df, provider, category, None))
    last = nav_df["date"].max()
    m["last_nav_date"] = last.date() if hasattr(last, "date") else last
    m["nav_age_days"] = (today - m["last_nav_date"]).days
    m["last_nav"] = float(nav_df.iloc[-1]["nav"])
    return m


def fetch_many(codes, refresh, workers, log=print):
    """{code: DataFrame or Exception} - failures are returned, never raised."""
    out = {}

    def one(code):
        try:
            return code, fetch_nav_data.get_nav_history(code, force_refresh=refresh)
        except Exception as exc:
            return code, exc

    codes = list(dict.fromkeys(codes))
    with ThreadPoolExecutor(max_workers=max(1, workers)) as ex:
        for i, (code, res) in enumerate(ex.map(one, codes), 1):
            out[code] = res
            if i % 25 == 0 or i == len(codes):
                log(f"    {i}/{len(codes)} NAV histories fetched")
    return out


def rank_peers(peer_metrics):
    """peer_metrics: list of metric dicts (with scheme_code).  -> {scheme_code: peer info}.
    Same scoring engine and weights.ini as shortlist.py, over the whole peer group - the hard
    filters are NOT applied, so you see where a fund really stands.  Metrics that this screener
    does not compute (portfolio quality, downside capture, quartile consistency) are simply
    missing and, under the default policy, the other weights are rescaled.  A fund whose
    measured weight is below the coverage floor is NOT ranked (it gets no peer entry), so
    one thin-data peer can neither crash nor distort the ranking."""
    from shortlist import score_and_rank
    rules = REVIEW_RULES
    df = pd.DataFrame(peer_metrics)
    for col in config.SCORE_WEIGHTS:
        if col not in df.columns:
            df[col] = None
    df = df[df["history_years"] >= rules["min_peer_history_years"]].reset_index(drop=True)
    if df.empty:
        return {}
    ranked = score_and_rank(df)
    ranked = ranked[ranked["status"] == "RANKED"].reset_index(drop=True)
    n = len(ranked)
    if n == 0:
        return {}

    def med(col):
        s = pd.to_numeric(ranked[col], errors="coerce").dropna()
        return float(s.median()) if len(s) else None

    stats = {"size": n, "median_3y": med("cagr_3y"), "median_5y": med("cagr_5y"),
             "median_dd": med("max_drawdown"), "unranked": len(df) - n}
    out = {}
    for _, r in ranked.iterrows():
        rank = int(r["rank"])
        pct = 100.0 * (n - rank) / (n - 1) if n > 1 else 50.0
        out[str(r["scheme_code"])] = dict(stats, rank=rank, percentile=pct)
    return out


def _pc(x, digits=1):
    return "n/a" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x * 100:.{digits}f}%"


def _pp(x):
    return f"{x * 100:+.1f} pp"


def decide(f, rules=REVIEW_RULES, rf=None):
    """Rules -> verdict.  f is a dict of facts about one holding (see analyse_holding).
    Returns {verdict, red, amber, green, notes, confidence, evidence}."""
    rf = config.RISK_FREE_RATE if rf is None else rf
    red, amber, green, notes = [], [], [], []
    level = {"absolute": 0, "benchmark": 0, "peers": 0}      # 0 none, 1 amber, 2 red - per evidence type

    def flag(kind, lvl, text):
        (red if lvl == 2 else amber).append(text)
        level[kind] = max(level[kind], lvl)

    def result(verdict, conf="-", ev="-"):
        return {"verdict": verdict, "red": red, "amber": amber, "green": green,
                "notes": notes, "confidence": conf, "evidence": ev}

    if f.get("zero_value"):
        red.append("Valuation is Rs 0 - the scheme appears to be wound up or closed. "
                   "Check with the AMC/RTA whether the units can still be redeemed.")
        return result(V_REVIEW, "n/a", "holdings file")
    if not f.get("mapped"):
        notes.append("Could not map this ISIN or name to an AMFI scheme - check the ISIN in the holdings file.")
        return result(V_NODATA)
    if f.get("nav_error"):
        notes.append(f"NAV history could not be fetched ({f['nav_error']}).")
        return result(V_NODATA)

    hard = False
    age = f.get("nav_age_days")
    if age is not None and age > rules["stale_nav_days"]:
        red.append(f"Last NAV is {age} days old ({f.get('last_nav_date')}) - the scheme may be "
                   f"closed, merged or renamed. Check with the AMC.")
        hard = True
    hist = f.get("history_years") or 0.0
    if hist < rules["min_history_years"] and not hard:
        notes.append(f"Only {hist:.1f} years of NAV history; a verdict needs {rules['min_history_years']:.0f}+. "
                     f"Judge it again when it is older.")
        return result(V_NODATA)

    used_bench = used_peers = False
    # -- absolute return
    c3, c5 = f.get("cagr_3y"), f.get("cagr_5y")
    if c3 is not None and c3 < rf:
        flag("absolute", 2, f"3-year CAGR {_pc(c3)} is below the {_pc(rf)} risk-free rate (a deposit would have done better).")
    elif c5 is not None and c5 < rf:
        flag("absolute", 1, f"5-year CAGR {_pc(c5)} is below the {_pc(rf)} risk-free rate.")
    # -- benchmark
    ex = f.get("benchmark_excess_return")
    if ex is not None:
        used_bench = True
        bn = f.get("benchmark_name") or "benchmark"
        if ex <= rules["bench_excess_red"]:
            flag("benchmark", 2, f"3-year return is {_pp(ex)} a year versus {bn}.")
        elif ex < rules["bench_excess_amber"]:
            flag("benchmark", 1, f"3-year return is {_pp(ex)} a year versus {bn}.")
        else:
            green.append(f"3-year return is {_pp(ex)} a year versus {bn}.")
    beat = f.get("rolling_beat_benchmark_pct")
    if beat is not None:
        used_bench = True
        txt = f"Beat the benchmark in {_pc(beat, 0)} of rolling 3-year windows."
        if beat < rules["rolling_beat_red"]:
            flag("benchmark", 2, txt)
        elif beat < rules["rolling_beat_amber"]:
            flag("benchmark", 1, txt)
        elif beat >= rules["rolling_beat_green"]:
            green.append(txt)
    # -- peers
    peer = f.get("peer")
    if peer and peer["size"] >= rules["min_peer_group_size"]:
        used_peers = True
        pct = peer["percentile"]
        txt = (f"Ranks {peer['rank']} of {peer['size']} in its category ({f.get('category_short')}) "
               f"- percentile {pct:.0f}, 100 = best.")
        if pct <= rules["peer_red_percentile"]:
            flag("peers", 2, txt)
        elif pct < rules["peer_amber_percentile"]:
            flag("peers", 1, txt)
        elif pct >= rules["peer_green_percentile"]:
            green.append(txt)
        dd, mdd = f.get("max_drawdown"), peer.get("median_dd")
        if dd is not None and mdd is not None and dd < mdd - rules["drawdown_worse_than_peers"]:
            flag("peers", 1, f"Worst fall {_pc(dd, 0)} is much deeper than the peer median {_pc(mdd, 0)}.")
    elif peer:
        notes.append(f"Only {peer['size']} peers with enough data - too few for a peer ranking.")
    elif f.get("peer_unranked"):
        notes.append("Not ranked against its peers: too little data to build a comparable score "
                     "(see the coverage floor in weights.ini).")

    if not used_peers and not used_bench:
        conf, ev = "Low", "absolute returns only"
        notes.append("No like-for-like peer group or benchmark for this category (index, fund-of-funds, "
                     "thematic or debt fund): judged on absolute returns only. Compare it with the index "
                     "or theme it tracks yourself.")
    elif used_peers and used_bench:
        conf, ev = "High", "benchmark + peers"
    else:
        conf, ev = "Medium", "peers only" if used_peers else "benchmark only"

    # Count evidence TYPES, not individual flags: excess return and rolling beat-rate both measure
    # the benchmark gap, so they must not count twice.
    r = sum(1 for v in level.values() if v == 2)
    a = sum(1 for v in level.values() if v == 1)
    if hard or r >= 2 or (r >= 1 and a >= 2):
        verdict = V_REVIEW
    elif r == 1 or a >= 2:
        verdict = V_WATCH
    else:
        verdict = V_CONTINUE
    return result(verdict, conf, ev)


# --------------------------------------------------------------------------
# Whole-portfolio analysis
# --------------------------------------------------------------------------
def short_category(cat):
    return (cat or "").split(" - ")[-1].strip() or "Unknown"


def analyse(rows, index, peers=True, refresh=False, workers=4, today=None, log=print):
    """Adds the analysis fields to each row dict (key 'A').  Returns run info for the Method sheet."""
    today = today or date.today()
    rules = REVIEW_RULES
    info = {"peer_groups": {}, "peers_failed": 0, "peer_capped": [], "peers_unranked": {}}

    # 1. map each distinct ISIN once
    units = {}                                    # isin|name -> shared analysis dict
    for r in rows:
        key = r["ISIN"] or r["Scheme Name"]
        if key in units:
            r["_u"] = units[key]
            continue
        rec, how = None, ""
        if r["ISIN"] in index.by_isin:
            rec, how = index.by_isin[r["ISIN"]], "ISIN"
        else:
            rec, ratio = index.match_name(r["Scheme Name"])
            if rec:
                how = f"Name match {ratio:.0%} (ISIN not found - verify)"
        zero = (_num(r.get("Valuation (Rs)")) or 0) == 0
        u = {"rec": rec, "how": how, "zero_value": zero, "mapped": rec is not None,
             "category": rec["category"] if rec else "", "isin": r["ISIN"]}
        u["category_short"] = short_category(u["category"])
        units[key] = u
        r["_u"] = u
    todo = [u for u in units.values() if u["mapped"] and not u["zero_value"]]
    log(f"Mapped {sum(u['mapped'] for u in units.values())} of {len(units)} distinct schemes to AMFI.")

    # 2. NAV history + metrics for the holdings
    log(f"Fetching NAV history for {len(todo)} schemes ...")
    navs = fetch_many([u["rec"]["code"] for u in todo], refresh, workers, log)
    bench = benchmark.BenchmarkProvider(force_refresh=refresh) if config.BENCHMARK_ENABLED else None
    for u in todo:
        res = navs[u["rec"]["code"]]
        if isinstance(res, Exception):
            u["nav_error"] = str(res)[:80]
            continue
        u["m"] = analyse_nav(res, u["category"], bench, today)

    # 3. peer ranking per category
    if peers:
        cats = {}
        for u in todo:
            if "m" in u and peer_eligible(u["category"]):
                cats.setdefault(u["category"], []).append(u)
        for cat, us in cats.items():
            own = {u["rec"]["code"] for u in us}
            pool = {}
            for r in index.records:
                if (r["open_ended"] and r["category"] == cat and is_direct_growth(r["name"])
                        and r["code"] not in pool):
                    pool[r["code"]] = r
            extra = sorted((c for c in pool if c not in own), key=lambda c: pool[c]["name"])
            room = max(0, config.MAX_FUNDS_TO_SCAN - len(own))
            if len(extra) > room:
                info["peer_capped"].append((short_category(cat), len(extra) + len(own)))
                extra = extra[:room]
            log(f"  peer group '{short_category(cat)}': {len(own)} of yours + {len(extra)} peers")
            got = fetch_many(extra, refresh, workers, lambda *_: None)
            metrics = [dict(u["m"], scheme_code=u["rec"]["code"]) for u in us]
            for code, res in got.items():
                if isinstance(res, Exception):
                    info["peers_failed"] += 1
                    continue
                try:
                    metrics.append(dict(analyse_nav(res, cat, bench, today), scheme_code=code))
                except Exception:
                    info["peers_failed"] += 1
            ranking = rank_peers(metrics)
            info["peer_groups"][short_category(cat)] = ranking[next(iter(ranking))]["size"] if ranking else 0
            if ranking:
                unranked = ranking[next(iter(ranking))]["unranked"]
                if unranked:
                    info["peers_unranked"][short_category(cat)] = unranked
            for u in us:
                u["peer"] = ranking.get(u["rec"]["code"])
                u["peer_unranked"] = bool(ranking) and u["peer"] is None
    if bench is not None:
        info["benchmarks"] = bench.describe()

    # 4. verdicts
    for u in units.values():
        m = u.get("m", {})
        facts = {"zero_value": u["zero_value"], "mapped": u["mapped"], "nav_error": u.get("nav_error"),
                 "category_short": u["category_short"], "peer": u.get("peer"),
                 "peer_unranked": u.get("peer_unranked", False)}
        facts.update({k: m.get(k) for k in (
            "history_years", "nav_age_days", "last_nav_date", "cagr_3y", "cagr_5y", "max_drawdown",
            "benchmark_excess_return", "rolling_beat_benchmark_pct", "benchmark_name")})
        u["d"] = decide(facts, rules)

    # 5. portfolio-level notes
    total = sum(_num(r.get("Valuation (Rs)")) or 0 for r in rows)
    by_cat, by_isin = {}, {}
    for r in rows:
        v = _num(r.get("Valuation (Rs)")) or 0
        u = r["_u"]
        c = u["category_short"] if u["mapped"] else "Unmapped"
        by_cat.setdefault(c, {"isins": set(), "value": 0.0})
        by_cat[c]["isins"].add(u["isin"] or r["Scheme Name"])
        by_cat[c]["value"] += v
        by_isin.setdefault(u["isin"] or r["Scheme Name"], {"value": 0.0, "folios": 0})
        by_isin[u["isin"] or r["Scheme Name"]]["value"] += v
        by_isin[u["isin"] or r["Scheme Name"]]["folios"] += 1
    ranks = {}                                    # category -> [(percentile, isin)]
    for u in units.values():
        if u.get("peer"):
            ranks.setdefault(u["category_short"], []).append((u["peer"]["percentile"], u["isin"]))
    for c in ranks:
        ranks[c].sort(reverse=True)
    for r in rows:
        u = r["_u"]
        key = u["isin"] or r["Scheme Name"]
        c = u["category_short"] if u["mapped"] else "Unmapped"
        v = _num(r.get("Valuation (Rs)")) or 0
        notes = []
        w_isin = by_isin[key]["value"] / total if total else 0
        if w_isin > rules["max_single_fund_weight"]:
            notes.append(f"Single fund is {w_isin:.0%} of your mutual fund portfolio (concentration).")
        n_cat = len(by_cat[c]["isins"])
        if u["mapped"] and peer_eligible(u["category"]) and n_cat >= rules["overlap_funds"]:
            line = (f"{n_cat} of your funds are in '{c}' ({by_cat[c]['value'] / total:.0%} of the portfolio) "
                    f"- likely overlap.") if total else ""
            order = [i for _, i in ranks.get(c, [])]
            if u["isin"] in order and len(order) >= 2:
                line += f" By peer rank this is #{order.index(u['isin']) + 1} of {len(order)} of yours."
            notes.append(line)
        if by_isin[key]["folios"] > 1:
            notes.append(f"Held in {by_isin[key]['folios']} folios (combined {w_isin:.1%} of the portfolio).")
        if u["mapped"] and "direct" not in u["rec"]["name"].lower():
            notes.append("AMFI lists this ISIN as a Regular-plan scheme: the Direct plan of the same fund "
                         "charges less. Check which plan you hold.")
        if u["mapped"] and any(w in u["rec"]["name"].lower() for w in ("idcw", "dividend")):
            notes.append("IDCW/dividend option: NAV-based returns understate what you received.")
        r["_notes"] = notes
        r["_weight"] = v / total if total else None
        r["_cat_count"], r["_cat_weight"] = n_cat, (by_cat[c]["value"] / total if total else None)
    info["total_value"] = total
    return info


# --------------------------------------------------------------------------
# Output columns
# --------------------------------------------------------------------------
KINDS = {"Closing Bal (Units)": "units", "NAV (Rs)": "nav", "Cumulative Amount Invested (INR)": "money",
         "Valuation (Rs)": "money", "Unrealised Profit/Loss": "money", "Unrealised Profit/Loss (%)": "num2",
         "PDF Page": "int"}
WIDTHS = {"Scheme Name": 52, "ISIN": 15, "Folio No.": 16, "Check": 30}

NEW_COLS = [  # (header, kind, width, wrap)
    ("Verdict", "text", 20, False), ("Confidence", "text", 11, False),
    ("Why (summary)", "text", 70, True), ("Red flags", "text", 60, True),
    ("Amber flags", "text", 60, True), ("Positives", "text", 60, True),
    ("Portfolio Notes", "text", 60, True), ("Data Notes", "text", 50, True),
    ("Portfolio Weight %", "pct", 11, False),
    ("AMFI Category", "text", 34, False), ("AMFI Scheme Code", "text", 11, False),
    ("Matched Scheme (AMFI)", "text", 52, False), ("Match Method", "text", 20, False),
    ("1Y CAGR", "pct", 9, False), ("3Y CAGR", "pct", 9, False), ("5Y CAGR", "pct", 9, False),
    ("Volatility", "pct", 9, False), ("Max Drawdown", "pct", 10, False),
    ("Sortino Ratio", "num2", 9, False), ("Rolling 3Y Positive %", "pct", 11, False),
    ("SIP XIRR (3Y)", "pct", 10, False),
    ("Benchmark", "text", 26, False), ("Benchmark 3Y CAGR", "pct", 11, False),
    ("3Y Excess vs Benchmark", "pct", 12, False), ("Rolling 3Y Beat Benchmark %", "pct", 13, False),
    ("Peer Group Size", "int", 9, False), ("Peer Rank", "int", 8, False),
    ("Peer Percentile (100=best)", "num0", 11, False),
    ("Peer Median 3Y CAGR", "pct", 11, False), ("Peer Median 5Y CAGR", "pct", 11, False),
    ("NAV History (yrs)", "num1", 9, False), ("Last NAV Date", "date", 12, False),
    ("Analysis Date", "date", 12, False),
]


def _bullets(items):
    return "\n".join(f"• {x}" for x in items)


def new_values(r, today):
    u = r["_u"]
    d, m = u["d"], u.get("m", {})
    peer = u.get("peer") or {}
    flags = d["red"] + d["amber"]
    summary = "; ".join((flags or d["green"])[:3]) or "; ".join(d["notes"][:1])
    data_notes = list(d["notes"])
    if u["mapped"] and u["how"] != "ISIN":
        data_notes.append("Matched by scheme name, not ISIN - verify the match.")
    vals = {
        "Verdict": d["verdict"], "Confidence": d["confidence"], "Why (summary)": summary,
        "Red flags": _bullets(d["red"]), "Amber flags": _bullets(d["amber"]),
        "Positives": _bullets(d["green"]), "Portfolio Notes": _bullets(r["_notes"]),
        "Data Notes": _bullets(data_notes), "Portfolio Weight %": r["_weight"],
        "AMFI Category": u["category"], "AMFI Scheme Code": u["rec"]["code"] if u["mapped"] else "",
        "Matched Scheme (AMFI)": u["rec"]["name"] if u["mapped"] else "", "Match Method": u["how"],
        "1Y CAGR": m.get("cagr_1y"), "3Y CAGR": m.get("cagr_3y"), "5Y CAGR": m.get("cagr_5y"),
        "Volatility": m.get("volatility"), "Max Drawdown": m.get("max_drawdown"),
        "Sortino Ratio": m.get("sortino_ratio"), "Rolling 3Y Positive %": m.get("rolling_positive_pct"),
        "SIP XIRR (3Y)": m.get("sip_xirr"), "Benchmark": m.get("benchmark_name") or "",
        "Benchmark 3Y CAGR": m.get("benchmark_cagr"), "3Y Excess vs Benchmark": m.get("benchmark_excess_return"),
        "Rolling 3Y Beat Benchmark %": m.get("rolling_beat_benchmark_pct"),
        "Peer Group Size": peer.get("size"), "Peer Rank": peer.get("rank"),
        "Peer Percentile (100=best)": peer.get("percentile"), "Peer Median 3Y CAGR": peer.get("median_3y"),
        "Peer Median 5Y CAGR": peer.get("median_5y"), "NAV History (yrs)": m.get("history_years"),
        "Last NAV Date": m.get("last_nav_date"), "Analysis Date": today,
    }
    return [vals[h] for h, *_ in NEW_COLS]


def build_sheets(headers, rows, info, today, source):
    all_headers = headers + [h for h, *_ in NEW_COLS]
    kinds = [KINDS.get(h, "text") for h in headers] + [k for _, k, _, _ in NEW_COLS]
    widths = [WIDTHS.get(h, 14) for h in headers] + [w for _, _, w, _ in NEW_COLS]
    wraps = [False] * len(headers) + [wr for *_, wr in NEW_COLS]
    verdict_col = len(headers)
    body = [[r.get(h) for h in headers] + new_values(r, today) for r in rows]
    tot = [None] * len(all_headers)
    tot[headers.index("Scheme Name")] = "TOTAL"
    for h in ("Cumulative Amount Invested (INR)", "Valuation (Rs)", "Unrealised Profit/Loss"):
        if h in headers:
            tot[headers.index(h)] = sum(_num(r.get(h)) or 0 for r in rows)
    main = {"name": "Holdings Analysis", "headers": all_headers, "rows": body + [tot], "kinds": kinds,
            "widths": widths, "wraps": wraps, "verdict_col": verdict_col, "freeze": (1, 2)}

    # Summary
    total = info["total_value"] or 0
    seen, per_v = set(), {}
    for r in rows:
        v = r["_u"]["d"]["verdict"]
        per_v.setdefault(v, {"isins": set(), "rows": 0, "value": 0.0})
        per_v[v]["isins"].add(r["ISIN"] or r["Scheme Name"])
        per_v[v]["rows"] += 1
        per_v[v]["value"] += _num(r.get("Valuation (Rs)")) or 0
    s = [[("Verdict summary", "title")] , ["Verdict", "Funds", "Folio rows", "Valuation (Rs)", "% of MF portfolio"]]
    for v in (V_REVIEW, V_WATCH, V_CONTINUE, V_NODATA):
        if v in per_v:
            p = per_v[v]
            s.append([(v, "verdict"), (len(p["isins"]), "int"), (p["rows"], "int"),
                      (p["value"], "money"), ((p["value"] / total) if total else None, "pct")])
    s += [[], [("Funds to look at first (REVIEW, then WATCH)", "title")],
          ["Scheme", "Verdict", "Valuation (Rs)", "Why"]]
    order = {V_REVIEW: 0, V_WATCH: 1}
    shown = set()
    for r in sorted(rows, key=lambda r: (order.get(r["_u"]["d"]["verdict"], 9), -(_num(r.get("Valuation (Rs)")) or 0))):
        d = r["_u"]["d"]
        key = (r["ISIN"], r["Scheme Name"])
        if d["verdict"] in order and key not in shown:
            shown.add(key)
            flags = d["red"] + d["amber"]
            s.append([r["Scheme Name"], (d["verdict"], "verdict"), (_num(r.get("Valuation (Rs)")), "money"),
                      "; ".join(flags[:3])])
    cats = {}
    for r in rows:
        c = r["_u"]["category_short"] if r["_u"]["mapped"] else "Unmapped"
        cats.setdefault(c, {"isins": set(), "value": 0.0})
        cats[c]["isins"].add(r["ISIN"] or r["Scheme Name"])
        cats[c]["value"] += _num(r.get("Valuation (Rs)")) or 0
    s += [[], [("Your portfolio by AMFI category", "title")], ["Category", "Funds", "Valuation (Rs)", "% of MF portfolio"]]
    for c, p in sorted(cats.items(), key=lambda kv: -kv[1]["value"]):
        s.append([c, (len(p["isins"]), "int"), (p["value"], "money"), ((p["value"] / total) if total else None, "pct")])
    summary = {"name": "Summary", "headers": None, "rows": s, "kinds": ["text"] * 5,
               "widths": [60, 18, 16, 80, 18], "wraps": [False, False, False, True, False]}

    # Method
    rl = REVIEW_RULES
    lines = [
        ("How to read this file", "title"),
        f"Analysis date: {today}.  Source holdings file: {source}.  Risk-free rate used: {config.RISK_FREE_RATE:.1%}.",
        "Each fund gets red / amber / green flags from three kinds of evidence, then a verdict:",
        "  The verdict counts evidence TYPES (absolute return, benchmark, peers), not individual flags, so two "
        "flags that measure the same thing never count twice.",
        "  REVIEW = red in 2 or more evidence types, or red in 1 and amber in 2, or a closed/stale scheme.  "
        "CONTINUE - WATCH = red in 1 type, or amber in 2.  CONTINUE = otherwise.  "
        "INSUFFICIENT DATA = not mapped, no NAV data, or less than "
        f"{rl['min_history_years']:.0f} years of history.",
        "Confidence: High = benchmark and peers both available; Medium = one of them; Low = absolute returns only.",
        "", ("Evidence and thresholds (edit REVIEW_RULES at the top of mf_holdings_screener.py)", "title"),
        f"1. Absolute: 3-year CAGR below the risk-free rate = red; 5-year below it = amber.",
        f"2. Benchmark (equity categories with a benchmark in config.py): 3-year excess return at or below "
        f"{rl['bench_excess_red']:.0%} = red, below {rl['bench_excess_amber']:.0%} = amber; beat the benchmark in "
        f"under {rl['rolling_beat_red']:.0%} of rolling 3-year windows = red, under {rl['rolling_beat_amber']:.0%} = amber.",
        f"3. Peers: rank among all Direct-Growth funds of the same AMFI category with {rl['min_peer_history_years']:.0f}+ "
        f"years of history, using shortlist.py's scoring engine and weights.ini "
        f"(weights file {os.path.basename(str(config.WEIGHTS_FILE))}, version {config.WEIGHTS_VERSION}; "
        f"policy {config.MISSING_DATA_POLICY}, method {config.SCORING_METHOD}). Portfolio quality, downside capture "
        f"and quartile consistency are not computed here, so they are left out and the other weights rescaled; a "
        f"fund whose measured weight is under the coverage floor ({config.MIN_SCORE_COVERAGE:.0%}) is not ranked. "
        f"Percentile at or below {rl['peer_red_percentile']} = red, below "
        f"{rl['peer_amber_percentile']} = amber. Needs at least {rl['min_peer_group_size']} ranked peers. "
        f"A max drawdown more than {rl['drawdown_worse_than_peers']:.0%} deeper than the peer median = amber.",
        f"4. A last NAV older than {rl['stale_nav_days']} days, or a Rs 0 valuation, means the scheme may be closed or merged = REVIEW.",
        f"Portfolio notes (do not change the verdict): a fund above {rl['max_single_fund_weight']:.0%} of the portfolio, "
        f"{rl['overlap_funds']}+ of your funds in one category, funds held in several folios, Regular-plan ISINs.",
        "", ("Limits - please read", "title"),
        "This is a rules-based screen of PAST NAV data. It is not investment advice and past returns do not guarantee future returns.",
        "It does not know your goals, holding period, tax position, exit load or the cost of switching. "
        "Selling can trigger tax and exit load; check current rules before acting on any verdict.",
        "Index funds, fund-of-funds, thematic/sectoral and debt funds have no like-for-like peer group here and are judged on "
        "absolute returns only (Confidence = Low). For index funds the things that matter - tracking error and expense "
        "ratio - are not in NAV data.",
        "Expense ratio, fund manager changes, AUM and portfolio holdings are not used.",
        "Benchmark excess return uses an index-fund proxy for the benchmark unless you supplied a TRI csv in "
        "data/benchmarks/ (see README), so it is slightly flattered.",
        "Peer groups come from AMFI's category for the scheme; a young fund with less history than needed is not ranked.",
        "", ("This run", "title"),
    ]
    for line in info.get("benchmarks", []):
        lines.append("Benchmark: " + line.strip())
    for g, n in info["peer_groups"].items():
        lines.append(f"Peer group '{g}': {n} funds ranked.")
    for g, n in info.get("peers_unranked", {}).items():
        lines.append(f"Peer group '{g}': {n} fund(s) left out of the ranking for too little data.")
    for g, n in info["peer_capped"]:
        lines.append(f"Peer group '{g}' had {n} funds; capped at config.MAX_FUNDS_TO_SCAN (alphabetical).")
    if info["peers_failed"]:
        lines.append(f"{info['peers_failed']} peer fund(s) could not be fetched and were left out.")
    method = {"name": "Method", "headers": None, "rows": [[x] for x in lines], "kinds": ["text"],
              "widths": [150], "wraps": [True]}
    return [main, summary, method]


# --------------------------------------------------------------------------
# Writers (.xls via xlwt, .xlsx via openpyxl)
# --------------------------------------------------------------------------
NUMFMT = {"pct": "0.0%", "num0": "0", "num1": "0.0", "num2": "0.00", "money": "#,##0.00", "units": "#,##0.000",
          "nav": "0.0000", "int": "0", "date": "DD-MMM-YYYY", "text": "General"}
VERDICT_COLOURS = {V_REVIEW: "FFC7CE", V_WATCH: "FFEB9C", V_CONTINUE: "C6EFCE", V_NODATA: "D9D9D9"}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        try:
            return v.item()
        except Exception:
            return v
    return v


def write_xls(path, sheets):
    import xlwt
    wb = xlwt.Workbook(encoding="utf-8")
    head = xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour gray25; align: wrap on, vert centre")
    title = xlwt.easyxf("font: bold on, height 260")
    colour = {V_REVIEW: "rose", V_WATCH: "light_yellow", V_CONTINUE: "light_green", V_NODATA: "gray25"}
    styles = {}

    def style(kind, wrap=False, fill=None):
        k = (kind, wrap, fill)
        if k not in styles:
            parts = []
            if wrap:
                parts.append("align: wrap on, vert top")
            if fill:
                parts.append(f"pattern: pattern solid, fore_colour {fill}")
            if kind in ("verdict",):
                parts.append("font: bold on")
            styles[k] = xlwt.easyxf(";".join(parts) if parts else "", num_format_str=NUMFMT.get(kind, "General"))
        return styles[k]

    for sh in sheets:
        ws = wb.add_sheet(sh["name"])
        r0 = 0
        if sh["headers"]:
            for j, h in enumerate(sh["headers"]):
                ws.write(0, j, h, head)
            ws.row(0).height_mismatch = True
            ws.row(0).height = 600
            r0 = 1
        for j, w in enumerate(sh["widths"]):
            ws.col(j).width = int(w * 256)
        for i, row in enumerate(sh["rows"]):
            for j, cell in enumerate(row):
                kind, val = sh["kinds"][min(j, len(sh["kinds"]) - 1)], cell
                if isinstance(cell, tuple):
                    val, kind = cell
                val = _clean(val)
                if val is None or val == "":
                    continue
                if kind == "title":
                    ws.write(r0 + i, j, val, title)
                    continue
                fill = None
                if kind == "verdict" or (sh.get("verdict_col") == j and isinstance(val, str)):
                    fill = colour.get(val)
                    kind = "verdict"
                if kind == "date" and isinstance(val, date) and not isinstance(val, datetime):
                    val = datetime(val.year, val.month, val.day)
                ws.write(r0 + i, j, val, style(kind, sh["wraps"][min(j, len(sh["wraps"]) - 1)], fill))
        if sh.get("freeze"):
            ws.set_panes_frozen(True)
            ws.set_horz_split_pos(sh["freeze"][0])
            ws.set_vert_split_pos(sh["freeze"][1])
    wb.save(path)


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh["name"])
        r0 = 1
        if sh["headers"]:
            ws.append(sh["headers"])
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="D9D9D9")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            r0 = 2
        for i, row in enumerate(sh["rows"]):
            for j, cell in enumerate(row):
                kind, val = sh["kinds"][min(j, len(sh["kinds"]) - 1)], cell
                if isinstance(cell, tuple):
                    val, kind = cell
                val = _clean(val)
                if val is None or val == "":
                    continue
                c = ws.cell(row=r0 + i, column=j + 1, value=val)
                if kind == "title":
                    c.font = Font(bold=True, size=13)
                    continue
                c.number_format = NUMFMT.get(kind, "General")
                if sh["wraps"][min(j, len(sh["wraps"]) - 1)]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                if kind == "verdict" or (sh.get("verdict_col") == j and isinstance(val, str)):
                    if val in VERDICT_COLOURS:
                        c.fill = PatternFill("solid", fgColor=VERDICT_COLOURS[val])
                    c.font = Font(bold=True)
        for j, w in enumerate(sh["widths"], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        if sh.get("freeze"):
            ws.freeze_panes = ws.cell(row=sh["freeze"][0] + 1, column=sh["freeze"][1] + 1)
    wb.save(path)


# --------------------------------------------------------------------------
def default_output(path):
    base, ext = os.path.splitext(path)
    return base + "-Screened" + (ext or ".XLS")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Screen each mutual fund in your holdings file and say whether to keep holding it.")
    ap.add_argument("holdings", help="mutual fund holdings .xls/.xlsx (from cdsl_holdings_to_xls.py)")
    ap.add_argument("-o", "--output", help="output file (.xls or .xlsx); default: <input name>-Screened.<ext>")
    ap.add_argument("--navall", help="AMFI NAVAll.txt saved on disk (otherwise it is downloaded)")
    ap.add_argument("--no-peers", action="store_true", help="skip the peer ranking (faster; fewer downloads)")
    ap.add_argument("--refresh", action="store_true", help="ignore cached NAV data and download again")
    ap.add_argument("--workers", type=int, default=4, help="parallel NAV downloads (default 4)")
    a = ap.parse_args(argv)

    out = a.output or default_output(a.holdings)
    headers, rows = read_holdings(a.holdings)
    print(f"Read {len(rows)} holdings from {a.holdings}")
    print("Loading AMFI scheme list (NAVAll.txt) ...")
    index = AmfiIndex(parse_navall(load_navall(a.navall, a.refresh)))
    if not index.records:
        sys.exit("AMFI NAVAll.txt had no scheme rows - is it the right file?")
    today = date.today()
    info = analyse(rows, index, peers=not a.no_peers, refresh=a.refresh, workers=a.workers, today=today)
    sheets = build_sheets(headers, rows, info, today, os.path.basename(a.holdings))
    (write_xlsx if out.lower().endswith(".xlsx") else write_xls)(out, sheets)

    print(f"\nWrote {out}\n")
    counts = {}
    for r in rows:
        counts.setdefault(r["_u"]["d"]["verdict"], []).append(r)
    for v in (V_REVIEW, V_WATCH, V_CONTINUE, V_NODATA):
        if v in counts:
            print(f"  {v}: {len(counts[v])} row(s)")
    for v in (V_REVIEW, V_WATCH):
        for r in counts.get(v, []):
            d = r["_u"]["d"]
            print(f"  [{v}] {r['Scheme Name'][:60]}: {'; '.join((d['red'] + d['amber'])[:2])}")
    if info.get("peer_capped"):
        print("\nNote: some peer groups were capped at config.MAX_FUNDS_TO_SCAN (see the Method sheet).")


if __name__ == "__main__":
    main()
