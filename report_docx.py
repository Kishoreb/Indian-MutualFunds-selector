"""
report_docx.py
--------------
Turns the Markdown text produced by check_fund.render() into a Word (.docx) file.

It understands only the small Markdown subset that render() emits:
    # / ## / ###   headings
    - item         bullet list
    | a | b |      tables (the |---| separator row is skipped; the first row is the header)
    **bold**       bold runs
    `code`         monospace runs
    _whole line_   italic paragraph
    anything else  normal paragraph

Needs:  pip install python-docx
"""

import re

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor
except ImportError:  # pragma: no cover
    Document = None

_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")
_SEPARATOR = re.compile(r"^\|\s*:?-{2,}")


class DocxUnavailable(RuntimeError):
    pass


def _add_runs(paragraph, text, bold=False, italic=False):
    """Adds text to a paragraph, honouring **bold** and `code` spans."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.italic = italic
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.bold = bold
            run.italic = italic
        else:
            run = paragraph.add_run(part)
            run.bold = bold
            run.italic = italic


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def _add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            _add_runs(para, text, bold=(i == 0))
            for run in para.runs:
                run.font.size = Pt(9)
            if i == 0:
                _shade(cell, "D9E2F3")
    doc.add_paragraph()


def markdown_to_docx(markdown_text, path, title=None):
    """Writes markdown_text to a .docx at `path`."""
    if Document is None:
        raise DocxUnavailable("python-docx is not installed. Run:  pip install python-docx")
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    if title:
        doc.core_properties.title = title

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _SEPARATOR.match(lines[i].strip()):
                    block.append(_split_row(lines[i]))
                i += 1
            if block:
                _add_table(doc, block)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading(level=level)
            _add_runs(h, m.group(2))
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, stripped[2:])
            i += 1
            continue
        if len(stripped) > 2 and stripped.startswith("_") and stripped.endswith("_"):
            p = doc.add_paragraph()
            _add_runs(p, stripped[1:-1], italic=True)
            i += 1
            continue
        p = doc.add_paragraph()
        _add_runs(p, stripped)
        i += 1

    doc.save(path)
    return path
